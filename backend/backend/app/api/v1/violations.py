"""
Violations Management Routes
"""
import os
import shutil
import json
from datetime import datetime
from uuid import UUID
from pathlib import Path
from io import BytesIO

from fastapi import APIRouter, Depends, HTTPException, status, Query, UploadFile, File, Form
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, and_

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.core.database import get_db
from app.core.dependencies import get_current_user, get_admin_user
from app.core.telegram_utils import send_violation_notification, send_verification_notification
from app.models import User, Violation, Camera, Notification
from app.schemas.violations import (
    ViolationCreate,
    ViolationResponse,
    ViolationListResponse,
    ViolationDetailResponse,
    ViolationStatusUpdate
)

router = APIRouter(prefix="/api/violations", tags=["Violations"])

# Create uploads directory
UPLOAD_DIR = Path("/app/uploads/violations")
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)


def save_upload_file(file: UploadFile) -> str:
    """Save uploaded file dan return path"""
    file_extension = Path(file.filename).suffix
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"violation_{timestamp}_{file.filename}"
    file_path = UPLOAD_DIR / filename
    
    with open(file_path, "wb") as f:
        shutil.copyfileobj(file.file, f)
    
    # Return relative path untuk database
    return f"/uploads/violations/{filename}"


def create_notifications_for_all_users(
    db: Session,
    violation_id: UUID,
    message: str
) -> None:
    """Create notification untuk semua users yang aktif"""
    users = db.query(User).filter(User.is_active == True).all()
    
    for user in users:
        notification = Notification(
            violation_id=violation_id,
            user_id=user.id,
            message=message,
            is_read=False
        )
        db.add(notification)
    
    db.commit()


# ========================
# POST - Create Violation (Edge AI)
# ========================
@router.post("", response_model=ViolationResponse, status_code=status.HTTP_201_CREATED)
async def create_violation(
    camera_id: str = Form(...),
    missing_apd: str = Form(...),  # JSON string
    confidence_score: float = Form(...),
    snapshot: UploadFile = File(...),
    db: Session = Depends(get_db)
):
    """
    Menerima data pelanggaran dan file snapshot dari Edge AI
    
    Akses: Edge AI / System Integration (no auth)
    
    - camera_id: UUID kamera (form)
    - missing_apd: JSON string {"helmet": True, "vest": False} (form)
    - confidence_score: Float 0-1 (form)
    - snapshot: Image file (binary)
    """
    try:
        camera_uuid = UUID(camera_id)
    except ValueError:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid camera ID format"
        )
    
    # Validate camera exists
    camera = db.query(Camera).filter(Camera.id == camera_uuid).first()
    if not camera:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Camera not found"
        )
    
    # Parse missing_apd JSON
    try:
        missing_apd_dict = json.loads(missing_apd)
        if not isinstance(missing_apd_dict, dict):
            raise ValueError("missing_apd must be a JSON object")
    except (json.JSONDecodeError, ValueError) as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid missing_apd JSON: {str(e)}"
        )
    
    # Validate confidence_score
    if not 0.0 <= confidence_score <= 1.0:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="confidence_score must be between 0.0 and 1.0"
        )
    
    # Validate file type
    allowed_extensions = {".jpg", ".jpeg", ".png", ".gif", ".bmp"}
    file_ext = Path(snapshot.filename).suffix.lower()
    if file_ext not in allowed_extensions:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"File type not allowed. Allowed: {', '.join(allowed_extensions)}"
        )
    
    # Save snapshot
    snapshot_path = save_upload_file(snapshot)
    
    # Create violation record
    new_violation = Violation(
        camera_id=camera_uuid,
        missing_apd=missing_apd_dict,
        confidence_score=confidence_score,
        snapshot_path=snapshot_path,
        status="Unverified"
    )
    
    db.add(new_violation)
    db.commit()
    db.refresh(new_violation)
    
    # Create notification untuk semua users
    missing_items = []
    for item, is_missing in missing_apd_dict.items():
        if is_missing:
            missing_items.append(f"❌ {item.upper()}")
    missing_text = ", ".join(missing_items) if missing_items else "All clear"
    
    notification_message = f"🚨 PELANGGARAN APD TERDETEKSI - Kamera: {camera.name} - APD Hilang: {missing_text} - Confidence: {confidence_score*100:.1f}%"
    create_notifications_for_all_users(db, new_violation.id, notification_message)
    
    # Send Telegram notification
    await send_violation_notification(
        camera_name=camera.name,
        missing_apd=missing_apd_dict,
        confidence_score=confidence_score,
        snapshot_url=None,  # Could add full URL if needed
        violation_timestamp=str(new_violation.timestamp)
    )
    
    return new_violation


# ========================
# GET - List Violations (Dashboard)
# ========================
@router.get("", response_model=list[ViolationListResponse])
async def get_violations(
    camera_id: UUID = Query(None, description="Filter by camera ID"),
    status_filter: str = Query(None, description="Filter by status: Unverified, Verified, False_Positive"),
    start_date: datetime = Query(None, description="Start date range"),
    end_date: datetime = Query(None, description="End date range"),
    limit: int = Query(50, ge=1, le=500, description="Max records (1-500)"),
    offset: int = Query(0, ge=0, description="Pagination offset"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Mengambil riwayat pelanggaran dengan pagination & filter
    
    Akses: Semua role
    
    Query Parameters:
    - camera_id: Filter by camera (optional)
    - status_filter: Filter by status (optional)
    - start_date: Start of date range (optional)
    - end_date: End of date range (optional)
    - limit: Max records (default 50)
    - offset: Pagination offset (default 0)
    """
    query = db.query(Violation)
    
    filters = []
    
    if camera_id:
        filters.append(Violation.camera_id == camera_id)
    
    if status_filter and status_filter in ["Unverified", "Verified", "False_Positive"]:
        filters.append(Violation.status == status_filter)
    
    if start_date:
        filters.append(Violation.timestamp >= start_date)
    
    if end_date:
        filters.append(Violation.timestamp <= end_date)
    
    if filters:
        query = query.filter(and_(*filters))
    
    # Order by timestamp descending
    query = query.order_by(Violation.timestamp.desc())
    
    # Pagination
    violations = query.offset(offset).limit(limit).all()
    
    return violations


# ========================
# GET - Violation Detail (Dashboard)
# ========================
@router.get("/{violation_id}", response_model=ViolationDetailResponse)
async def get_violation_detail(
    violation_id: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Mengambil detail satu pelanggaran spesifik
    
    Akses: Semua role
    
    Return termasuk:
    - Snapshot path (untuk display foto)
    - All APD details
    - Confidence score
    """
    try:
        violation_uuid = UUID(violation_id)
    except ValueError:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid violation ID format"
        )
    
    violation = db.query(Violation).filter(Violation.id == violation_uuid).first()
    if not violation:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Violation not found"
        )
    
    return violation


# ========================
# PUT - Update Violation Status (Dashboard)
# ========================
@router.put("/{violation_id}/status", response_model=ViolationDetailResponse)
async def update_violation_status(
    violation_id: str,
    status_update: ViolationStatusUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Mengubah status pelanggaran menjadi Verified atau False Positive
    
    Akses: Semua role (yang bisa verify)
    
    Status yang valid:
    - Verified: Pelanggaran terkonfirmasi
    - False_Positive: Bukan pelanggaran
    """
    try:
        violation_uuid = UUID(violation_id)
    except ValueError:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Invalid violation ID format"
        )
    
    violation = db.query(Violation).filter(Violation.id == violation_uuid).first()
    if not violation:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Violation not found"
        )
    
    # Update status
    violation.status = status_update.status
    violation.notes = status_update.notes  # Store verification notes
    db.commit()
    db.refresh(violation)
    
    # Create notification untuk semua users
    if status_update.status == "Verified":
        notification_message = f"✅ PELANGGARAN TERVERIFIKASI - Kamera: {camera.name} - Diverifikasi oleh: {current_user.username} - Catatan: {status_update.notes or 'N/A'}"
    else:
        notification_message = f"⛔ BUKAN PELANGGARAN - Kamera: {camera.name} - Diverifikasi oleh: {current_user.username} - Catatan: {status_update.notes or 'N/A'}"
    
    create_notifications_for_all_users(db, violation.id, notification_message)
    
    # Send Telegram notification
    camera = db.query(Camera).filter(Camera.id == violation.camera_id).first()
    if camera:
        await send_verification_notification(
            camera_name=camera.name,
            status=status_update.status,
            verified_by=current_user.username,
            notes=status_update.notes
        )
    
    return violation


# ========================
# GET - Export Violations to Excel
# ========================
@router.get("/export/xlsx")
async def export_violations_xlsx(
    camera_id: UUID = Query(None, description="Filter by camera ID"),
    status_filter: str = Query(None, description="Filter by status"),
    start_date: datetime = Query(None, description="Start date"),
    end_date: datetime = Query(None, description="End date"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    """
    Mengunduh file .xlsx berisi riwayat pelanggaran
    
    Akses: Semua role
    
    Export fields:
    - ID, Camera ID, Timestamp, APD Missing
    - Confidence Score, Status, Snapshot Path
    """
    query = db.query(Violation)
    
    filters = []
    
    if camera_id:
        filters.append(Violation.camera_id == camera_id)
    
    if status_filter and status_filter in ["Unverified", "Verified", "False_Positive"]:
        filters.append(Violation.status == status_filter)
    
    if start_date:
        filters.append(Violation.timestamp >= start_date)
    
    if end_date:
        filters.append(Violation.timestamp <= end_date)
    
    if filters:
        query = query.filter(and_(*filters))
    
    violations = query.order_by(Violation.timestamp.desc()).all()
    
    if not violations:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="No violations found for export"
        )
    
    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Violations"
    
    # Headers
    headers = ["ID", "Camera ID", "Timestamp", "Missing APD", "Confidence Score", "Status", "Snapshot Path"]
    ws.append(headers)
    
    # Style headers
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Add data
    for violation in violations:
        missing_apd_str = ", ".join([k for k, v in violation.missing_apd.items() if v])
        ws.append([
            str(violation.id),
            str(violation.camera_id),
            violation.timestamp.isoformat(),
            missing_apd_str,
            f"{violation.confidence_score:.1%}",
            violation.status,
            violation.snapshot_path
        ])
    
    # Adjust column widths
    ws.column_dimensions['A'].width = 36
    ws.column_dimensions['B'].width = 36
    ws.column_dimensions['C'].width = 20
    ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 35
    
    # Save to BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)
    
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=violations_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"}
    )
